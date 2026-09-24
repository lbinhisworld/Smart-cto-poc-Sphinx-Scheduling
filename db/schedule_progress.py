"""已排 / 未排三类、待排负荷。统计不触发排程，不写交期。"""

from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from db.plan_facts import counts_as_pending_load, classify_order, export_label
from db.plan_store import current_plan_version
from db.prod_stats_seed import is_capacity_fixture_order
from db.dispatch_export import group_label
from db.snapshot import header_order_no, load_schedule_input
from db.tables import MdItemRow, PendingRollRow, SoOrderRow, WoRow, WoTaskRow
from engine.bom import gross_board_for_line
from engine.capacity import hours_man, hours_wall
from engine.expand import expand_order
from engine.kit_pool import KitSnapshot
from engine.models import ComponentRole

HOURS_LABEL = "需求工时，未落到日期"


def schedule_progress(session: Session) -> dict:
    version = current_plan_version(session)
    orders = [
        row
        for row in session.scalars(select(SoOrderRow).order_by(SoOrderRow.due_date, SoOrderRow.order_no)).all()
        if not is_capacity_fixture_order(row.order_no, row.order_source)
    ]
    wos = []
    tasks = []
    if version > 0:
        wos = list(session.scalars(select(WoRow).where(WoRow.plan_version == version)).all())
        tasks = list(session.scalars(select(WoTaskRow).where(WoTaskRow.plan_version == version)).all())
    rolls = list(
        session.scalars(select(PendingRollRow).where(PendingRollRow.status == "PENDING_CONFIRMATION")).all()
    )
    tasks_by_wo: dict[str, list[WoTaskRow]] = defaultdict(list)
    for t in tasks:
        tasks_by_wo[t.wo_no].append(t)
    wos_by_order: dict[str, list[WoRow]] = defaultdict(list)
    for w in wos:
        wos_by_order[header_order_no(w.source_order_no)].append(w)
    roll_orders = {header_order_no(r.source_order_no) for r in rolls}

    buckets: dict[str, list[dict]] = {
        "PLACED": [],
        "NOT_PLACED": [],
        "WIP_UNPLANNED": [],
        "PARTIAL": [],
        "POOL_OPEN": [],
    }
    for order in orders:
        own = wos_by_order.get(order.order_no, [])
        own_tasks = [t for w in own for t in tasks_by_wo.get(w.wo_no, [])]
        has_tasks = bool(own_tasks)
        has_unplaced = any(
            sum(t.qty_board for t in tasks_by_wo.get(w.wo_no, [])) < int(w.qty_board_plan) for w in own
        )
        has_completion = any(int(w.qty_board_done or 0) > 0 for w in own)
        has_roll = order.order_no in roll_orders
        bucket = classify_order(
            phase=order.schedule_phase or "PENDING",
            has_tasks=has_tasks,
            has_unplaced=has_unplaced,
            has_completion=has_completion,
            has_pending_roll=has_roll,
        )
        wall = sum((Decimal(str(t.hours_wall)) for t in own_tasks), Decimal(0))
        man = sum((Decimal(str(t.hours_man)) for t in own_tasks), Decimal(0))
        starts = [w.plan_start for w in own if w.plan_start]
        ends = [w.plan_end for w in own if w.plan_end]
        buckets[bucket].append(
            {
                "order_no": order.order_no,
                "customer": order.customer,
                "item_code": order.item_code,
                "due_date": order.due_date.isoformat(),
                "plan_start": min(starts).isoformat() if starts else None,
                "plan_end": max(ends).isoformat() if ends else None,
                "released": bool(own) and all(w.status == "RELEASED" for w in own),
                "hours_wall": _hours(wall) if has_tasks else None,
                "hours_man": _hours(man) if has_tasks else None,
                "bucket": bucket,
                "export_label": export_label(bucket, has_tasks=has_tasks),
                "note": "已在排程池，本轮未跑完" if bucket == "POOL_OPEN" else "",
                "has_tasks": has_tasks,
            }
        )
    counts = {key.lower(): len(rows) for key, rows in buckets.items()}
    return {
        "plan_version": version,
        "counts": {
            "placed": counts["placed"],
            "not_placed": counts["not_placed"],
            "wip_unplanned": counts["wip_unplanned"],
            "partial": counts["partial"],
            "pool_open": counts["pool_open"],
        },
        "placed": buckets["PLACED"],
        "not_placed": buckets["NOT_PLACED"],
        "wip_unplanned": buckets["WIP_UNPLANNED"],
        "partial": buckets["PARTIAL"],
        "pool_open": buckets["POOL_OPEN"],
    }


def pending_load(session: Session, *, today: date) -> dict:
    progress = schedule_progress(session)
    headers = [row["order_no"] for row in progress["not_placed"]]
    if not headers:
        return _empty_load()
    inp = load_schedule_input(session, today=today, order_nos=headers)
    names = {
        r.item_code: r.item_name
        for r in session.scalars(select(MdItemRow)).all()
    }
    orders = sorted(inp.orders, key=lambda o: (o.due_date, o.order_no))
    header_set = set(headers)
    orders = [o for o in orders if header_order_no(o.order_no) in header_set]
    pool = KitSnapshot(dict(inp.stock))
    gross_by: dict[str, int] = defaultdict(int)
    shortage_by: dict[str, int] = defaultdict(int)
    semi_net: dict[str, int] = defaultdict(int)
    role_of: dict[str, str] = {}
    contributors: dict[str, list[dict]] = defaultdict(list)
    uncomputable: list[str] = []
    by_order: list[dict] = []
    by_group: dict[str, dict[str, Decimal]] = defaultdict(lambda: {"wall": Decimal(0), "man": Decimal(0)})
    total_wall = Decimal(0)
    total_man = Decimal(0)

    for order in orders:
        item = inp.items.get(order.item_code)
        if item is None or not item.computable:
            uncomputable.append(header_order_no(order.order_no))
            continue
        try:
            sph = inp.sph_of(order.item_code, item.group_code)
        except KeyError:
            uncomputable.append(header_order_no(order.order_no))
            continue
        wo = expand_order(order, item, inp.converts_for(order.item_code), today, sph=sph)
        if wo is None:
            uncomputable.append(header_order_no(order.order_no))
            continue
        wall = hours_wall(wo.qty_board_plan, sph, wo.crew_plan, inp.converts_for(order.item_code))
        man = hours_man(wall, wo.crew_plan)
        semi_wall = Decimal(0)
        semi_man = Decimal(0)
        lines = inp.bom_for(order.item_code)
        for line in lines:
            gross = gross_board_for_line(order, line, item)
            _from, net = pool.consume(
                order_no=order.order_no,
                component_item_code=line.component_item_code,
                gross_board=gross,
                line_no=line.line_no,
            )
            gross_by[line.component_item_code] += gross
            role_of[line.component_item_code] = line.component_role.value
            if line.component_role == ComponentRole.PURCHASED:
                shortage_by[line.component_item_code] += net
            if line.component_role == ComponentRole.SEMI:
                semi_net[line.component_item_code] += net
            contributors[line.component_item_code].append(
                {"order_no": header_order_no(order.order_no), "gross_board": gross}
            )
            if line.component_role == ComponentRole.SEMI and net > 0:
                semi_item = inp.items.get(line.component_item_code)
                if semi_item is not None:
                    try:
                        semi_sph = inp.sph_of(semi_item.item_code, semi_item.group_code)
                    except KeyError:
                        semi_sph = None
                    if semi_sph is not None:
                        sw = hours_wall(
                            net, semi_sph, semi_sph.crew_std, inp.converts_for(semi_item.item_code)
                        )
                        sm = hours_man(sw, semi_sph.crew_std)
                        semi_wall += sw
                        semi_man += sm
                        gkey = f"{semi_item.dept.value}|{semi_item.group_code.value}"
                        by_group[gkey]["wall"] += sw
                        by_group[gkey]["man"] += sm
        gkey = f"{item.dept.value}|{item.group_code.value}"
        by_group[gkey]["wall"] += wall
        by_group[gkey]["man"] += man
        total_wall += wall + semi_wall
        total_man += man + semi_man
        by_order.append(
            {
                "order_no": header_order_no(order.order_no),
                "finished_wall": _hours(wall),
                "finished_man": _hours(man),
                "semi_wall": _hours(semi_wall),
                "semi_man": _hours(semi_man),
            }
        )

    components = []
    for code, gross in sorted(gross_by.items()):
        components.append(
            {
                "component_item_code": code,
                "component_name": names.get(code, code),
                "role": "半成品" if role_of.get(code) == "SEMI" else "外购",
                "gross_board": gross,
                "shortage_board": int(shortage_by.get(code, 0)),
                "net_board": int(semi_net.get(code, 0)),
                "orders": contributors[code],
            }
        )
    return {
        "label": HOURS_LABEL,
        "order_count": len(headers),
        "uncomputable_count": len(set(uncomputable)),
        "uncomputable_orders": sorted(set(uncomputable)),
        "components": components,
        "hours": {
            "wall": _hours(total_wall),
            "man": _hours(total_man),
            "by_group": [
                {
                    "group": group_label(*key.split("|", 1)),
                    "wall": _hours(v["wall"]),
                    "man": _hours(v["man"]),
                }
                for key, v in sorted(by_group.items())
            ],
            "by_order": by_order,
        },
    }


def unscheduled_workbook_bytes(session: Session, *, today: date) -> bytes:
    progress = schedule_progress(session)
    load = pending_load(session, today=today)
    wb = Workbook()
    sheets = [
        ("已下单未排", progress["not_placed"] + [r for r in progress["pool_open"] if not r["has_tasks"]]),
        ("在制在途未排", progress["wip_unplanned"]),
        ("只排了一部分", progress["partial"] + [r for r in progress["pool_open"] if r["has_tasks"]]),
    ]
    first = True
    for title, rows in sheets:
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        first = False
        ws.append(["订单号", "客户", "品项", "交期", "计划开工", "计划完工", "已发布", "墙钟工时", "人时", "说明"])
        for row in rows:
            ws.append(
                [
                    row["order_no"],
                    row["customer"],
                    row["item_code"],
                    row["due_date"],
                    row["plan_start"] or "",
                    row["plan_end"] or "",
                    "是" if row["released"] else "",
                    row["hours_wall"] or "",
                    row["hours_man"] or "",
                    row["note"],
                ]
            )
        if not rows:
            ws.append(["（无）"])
    pick = wb.create_sheet("待排毛需求")
    pick.append([HOURS_LABEL, f"墙钟 {load['hours']['wall']}", f"人·时 {load['hours']['man']}"])
    pick.append(["子件编码", "子件名称", "角色", "毛需求(版)", "外购缺口(版)", "半成品净需求(版)"])
    for comp in load["components"]:
        pick.append(
            [
                comp["component_item_code"],
                comp["component_name"],
                comp["role"],
                comp["gross_board"],
                comp["shortage_board"],
                comp["net_board"],
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hours(value: Decimal) -> str:
    return format(value.quantize(Decimal("0.01")), "f")


def _empty_load() -> dict:
    return {
        "label": HOURS_LABEL,
        "order_count": 0,
        "uncomputable_count": 0,
        "uncomputable_orders": [],
        "components": [],
        "hours": {"wall": "0.00", "man": "0.00", "by_group": [], "by_order": []},
    }

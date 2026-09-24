"""Phase 5：派工单 Excel 导出。"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from db.dispatch_export import build_dispatch_workbook_bytes
from db.repositories import run_schedule
from db.session import session_factory
from db.seed import import_seed_json, needs_seed_reload, reload_seed_json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SEED = ROOT / "seed" / "seed_data.json"
DB = ROOT / "data" / "scheduling_test_export.db"


@pytest.fixture()
def db_session():
    from db.session import init_db, make_engine

    if DB.is_file():
        DB.unlink()
    engine = make_engine(DB)
    init_db(engine)
    from db.migrate import ensure_schema

    ensure_schema(engine)
    factory = session_factory(engine)
    session = factory()
    if needs_seed_reload(session, SEED):
        reload_seed_json(session, SEED)
    else:
        import_seed_json(session, SEED)
    session.commit()
    yield session
    session.close()
    if DB.is_file():
        DB.unlink()


def test_dispatch_export_has_task_rows(db_session):
    run_schedule(
        db_session,
        today=date(2026, 9, 15),
        order_nos=["SO-001", "SO-002", "SO-003"],
        persist=True,
        trigger="export-test",
    )
    db_session.commit()
    raw = build_dispatch_workbook_bytes(db_session)
    wb = load_workbook(BytesIO(raw))
    assert wb.sheetnames == ["生产排程"]
    ws = wb["生产排程"]
    grid = " ".join(str(c.value or "") for row in ws.iter_rows() for c in row)
    assert "计划版本" not in grid
    assert "未发布" not in grid
    footer = ws.oddFooter.left.text or ""
    assert "未发布，不得下发" in footer
    assert "计划版本" in footer
    header_row = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1) if row[0] == "生产人员")
    headers = [c.value for c in ws[header_row]][:21]
    assert headers == [
        "生产人员",
        "班别",
        "型号",
        "品名",
        "枚/版",
        "单位",
        "订单需求量",
        "单据编号",
        "购货单位",
        "巧克力颜色",
        "气泡垫",
        "真空袋/内衬",
        "回料袋颜色",
        "单位",
        "SPH",
        "人力/人",
        "计划盒数",
        "实际生产",
        "生产工时",
        "累计工时",
        "备注",
    ]
    title = ws.cell(header_row - 2, 1).value
    assert title.startswith("《") and title.endswith("》生产排程")
    assert "·" not in title
    assert ws.cell(header_row - 1, 1).value == "组长："
    assert ws.cell(header_row - 1, 2).value in (None, "")
    assert str(ws.cell(header_row - 1, 16).value).startswith("生产日期：")
    assert ws.cell(header_row, 2).fill.fgColor.rgb.endswith("FFC000")
    assert ws.cell(header_row, 10).fill.fgColor.rgb.endswith("92D050")
    assert ws.cell(header_row, 15).fill.fgColor.rgb.endswith("00B0F0")
    assert ws.cell(header_row, 18).fill.fgColor.rgb.endswith("92D050")
    rows = [
        r
        for r in ws.iter_rows(min_row=header_row + 1, values_only=True)
        if r[2] and r[0] != "生产人员"
    ]
    assert len(rows) >= 3
    assert all(r[0] in (None, "") for r in rows)
    assert all(r[1] and "MANUAL" not in str(r[1]) and "MOLD" not in str(r[1]) for r in rows)
    assert all(str(r[7]).startswith("SO-") for r in rows)


def test_released_dispatch_has_no_preview_banner(db_session):
    from sqlalchemy import update

    from db.dispatch_export import build_dispatch_workbook
    from db.tables import WoRow

    run_schedule(
        db_session,
        today=date(2026, 9, 15),
        order_nos=["SO-001"],
        persist=True,
        trigger="export-release",
    )
    db_session.execute(update(WoRow).values(status="RELEASED"))
    db_session.commit()
    raw, filename = build_dispatch_workbook(db_session)
    assert filename == "dispatch.xlsx"
    wb = load_workbook(BytesIO(raw))
    ws = wb["生产排程"]
    grid = " ".join(str(c.value or "") for row in ws.iter_rows() for c in row)
    footer = ws.oddFooter.left.text or ""
    assert "未发布" not in grid
    assert "未发布" not in footer
    assert "计划版本" in footer
    assert "计划版本" not in grid


def test_earliest_plan_keeps_due_and_version(db_session):
    from db.earliest_plan import TITLE, build_earliest_plan
    from db.plan_store import current_plan_version
    from db.tables import SoOrderRow

    run_schedule(
        db_session,
        today=date(2026, 9, 15),
        order_nos=["SO-004"],
        persist=True,
        trigger="earliest-test",
    )
    db_session.commit()
    due = db_session.get(SoOrderRow, "SO-004").due_date
    version = current_plan_version(db_session)
    plan = build_earliest_plan(db_session, order_no="SO-004", today=date(2026, 9, 15))
    db_session.commit()
    assert db_session.get(SoOrderRow, "SO-004").due_date == due
    assert current_plan_version(db_session) == version
    assert plan["original_due"] == due.isoformat()
    if plan["eligible"]:
        assert plan["title"] == TITLE
        assert plan["suggested_due"] > due.isoformat()
        assert "未下发" in plan["title"]
        if plan["deliverable"]:
            assert len(plan["sales_sentences"]) == 3
        else:
            assert "仍未排完" in plan["note"]
    else:
        assert plan["reason"]


def test_pending_load_skips_orders_already_in_plan(db_session):
    from db.schedule_progress import pending_load, schedule_progress

    run_schedule(
        db_session,
        today=date(2026, 9, 15),
        order_nos=["SO-001"],
        persist=True,
        trigger="progress-test",
    )
    db_session.commit()
    progress = schedule_progress(db_session)
    not_placed = {r["order_no"] for r in progress["not_placed"]}
    assert "SO-001" not in not_placed
    load = pending_load(db_session, today=date(2026, 9, 15))
    assert load["label"] == "需求工时，未落到日期"
    assert "SO-001" not in {r["order_no"] for r in load["hours"]["by_order"]}
    assert "pool_open" in progress["counts"]


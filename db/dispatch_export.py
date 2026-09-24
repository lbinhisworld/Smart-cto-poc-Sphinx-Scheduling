"""派工单 Excel：客户《组》生产排程一张表。不写订单交期。"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.orm import Session

from db.bom_view import _bom_lines_for, _load_item, _route_for
from db.labor_cost_queries import WORK_CENTERS
from db.plan_facts import DISPATCH_HEADERS
from db.plan_store import current_plan_version
from db.tables import MdItemRow, MdSphRow, SoOrderRow, WoRow, WoTaskRow
from engine.bom import gross_board_for_line
from engine.expand import gross_semi_board
from engine.models import ComponentRole, Dept, GroupCode, Order, Uom, Wo, WoStatus, WoType

# 与客户表同一列序。两个「单位」底色不同，按位置取色，不能按列名。
_HEADER_COLORS = (
    None,
    "FFC000",
    "FFC000",
    "FFC000",
    "FFC000",
    "FFC000",
    "FFC000",
    "FFC000",
    "FFC000",
    "92D050",
    "92D050",
    "92D050",
    "92D050",
    "00B0F0",
    "00B0F0",
    "00B0F0",
    "00B0F0",
    "92D050",
    "00B0F0",
    "00B0F0",
    None,
)
_THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


_UOM_LABEL = {
    "PCS": "枚",
    "BOARD": "版",
    "BOX": "盒",
    "PACK": "包",
    "BAG": "袋",
    "CARTON": "箱",
}


def form_title(label: str) -> str:
    short = label.split("·")[-1].strip()
    if not short:
        return "《生产排程》"
    return f"《{short}》生产排程"


def group_label(dept: str, group_code: str) -> str:
    for d, g, name in WORK_CENTERS:
        if d == dept and g == group_code:
            return name
    return group_code


def role_label(role: str) -> str:
    if role == ComponentRole.SEMI.value or role == "SEMI":
        return "半成品（自做或领库存）"
    return "外购（领用）"


def build_dispatch_workbook(
    session: Session,
    *,
    group: str | None = None,
    dept: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> tuple[bytes, str]:
    version = current_plan_version(session)
    if version <= 0:
        wb = Workbook()
        ws = wb.active
        ws.title = "组排程"
        ws.append(["请先倒排"])
        return _save(wb), "dispatch.xlsx"

    group_code, dept_code = _resolve_group_filter(group, dept)
    tasks = list(
        session.scalars(
            select(WoTaskRow)
            .where(WoTaskRow.plan_version == version)
            .order_by(WoTaskRow.task_date, WoTaskRow.group_code, WoTaskRow.seq, WoTaskRow.task_id)
        ).all()
    )
    tasks = [
        t
        for t in tasks
        if _task_in_range(t, group_code=group_code, dept_code=dept_code, date_from=date_from, date_to=date_to)
    ]
    wo_nos = {t.wo_no for t in tasks}
    wos = {
        w.wo_no: w
        for w in session.scalars(select(WoRow).where(WoRow.wo_no.in_(wo_nos or {"-"}))).all()
    }
    preview = not tasks or any(wos[t.wo_no].status != "RELEASED" for t in tasks if t.wo_no in wos)
    note = f"计划版本 {version}" + (" · 未发布，不得下发" if preview else "")

    item_codes = {w.item_code for w in wos.values()}
    items = {
        i.item_code: i
        for i in session.scalars(select(MdItemRow).where(MdItemRow.item_code.in_(item_codes or {"-"}))).all()
    }
    orders = {
        o.order_no: o
        for o in session.scalars(
            select(SoOrderRow).where(SoOrderRow.order_no.in_({w.source_order_no.split("#L")[0] for w in wos.values()} or {"-"}))
        ).all()
    }
    sph_rows = {
        (s.item_code, s.group_code): s
        for s in session.scalars(select(MdSphRow).where(MdSphRow.item_code.in_(item_codes or {"-"}))).all()
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "生产排程"
    ws.oddFooter.left.text = note
    ws.evenFooter.left.text = note
    blocks = _blocks(tasks, wos)
    row = 1
    if not blocks:
        _write_block(ws, row, title="《生产排程》", work_date=None, rows=[["该范围没有任务"]])
    for (label, work_date), block_tasks in blocks:
        rows = [
            _dispatch_row(session, t, wos, items, orders, sph_rows)
            for t in block_tasks
            if t.wo_no in wos
        ]
        row = _write_block(ws, row, title=form_title(label), work_date=work_date, rows=rows) + 2

    filename = "dispatch-preview.xlsx" if preview else "dispatch.xlsx"
    return _save(wb), filename


def build_dispatch_workbook_bytes(session: Session) -> bytes:
    data, _name = build_dispatch_workbook(session)
    return data


def _save(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _resolve_group_filter(group: str | None, dept: str | None) -> tuple[str | None, str | None]:
    if not group:
        return None, dept
    text = group.strip()
    for d, g, name in WORK_CENTERS:
        if text in {g, name} and (dept is None or dept == d):
            return g, dept or d
    return text, dept


def _task_in_range(task: WoTaskRow, *, group_code, dept_code, date_from, date_to) -> bool:
    if group_code and task.group_code != group_code:
        return False
    if dept_code and (task.dept or "") != dept_code:
        return False
    if date_from and task.task_date < date_from:
        return False
    if date_to and task.task_date > date_to:
        return False
    return True


def _write_block(ws, row_no: int, *, title: str, work_date: date | None, rows: list[list]) -> int:
    width = len(DISPATCH_HEADERS)
    title_cell = ws.cell(row_no, 1, title)
    ws.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=width)
    title_cell.font = Font(name="宋体", size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    row_no += 1
    ws.cell(row_no, 1, "组长：")
    date_cell = ws.cell(row_no, 16, "生产日期：" + (_sheet_date(work_date) if work_date else ""))
    date_cell.alignment = Alignment(horizontal="right", vertical="center")
    row_no += 1
    for col, (name, color) in enumerate(zip(DISPATCH_HEADERS, _HEADER_COLORS), start=1):
        cell = ws.cell(row_no, col, name)
        cell.font = Font(name="宋体", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN
        if color:
            cell.fill = PatternFill("solid", fgColor=color)
    row_no += 1
    if len(rows) == 1 and rows[0] == ["该范围没有任务"]:
        ws.cell(row_no, width, "该范围没有任务")
        return row_no
    total = Decimal("0")
    for values, hours in rows:
        total += hours
        values[18] = _hours(hours)
        values[19] = _hours(total)
        for col, value in enumerate(values, start=1):
            ws.cell(row_no, col, value)
        row_no += 1
    return row_no - 1


def _blocks(tasks: list[WoTaskRow], wos: dict[str, WoRow]) -> list[tuple[tuple[str, date], list[WoTaskRow]]]:
    grouped: dict[tuple[str, date], list[WoTaskRow]] = {}
    for task in tasks:
        wo = wos.get(task.wo_no)
        if wo is None:
            continue
        label = group_label(task.dept or wo.dept, task.group_code)
        grouped.setdefault((label, task.task_date), []).append(task)
    for block in grouped.values():
        block.sort(key=lambda task: (task.seq, task.task_id))
    return sorted(grouped.items(), key=lambda item: (item[0][1], item[0][0]))


def _dispatch_row(session, task: WoTaskRow, wos, items, orders, sph_rows) -> tuple[list, Decimal]:
    wo = wos[task.wo_no]
    item = items.get(wo.item_code)
    header = wo.source_order_no.split("#L")[0]
    order = orders.get(header)
    names = [c["name"] for c in _finished_components(session, wo, orders) if c["name"] and c["name"] != "无下级子件"]
    sph = sph_rows.get((wo.item_code, task.group_code))
    hours = Decimal(str(task.hours_wall))
    board_per_box = Decimal(str(item.board_per_box)) if item is not None else Decimal(0)
    row = [
        "",
        group_label(task.dept or wo.dept, task.group_code),
        wo.item_code,
        item.item_name if item is not None else "",
        item.pcs_per_board if item is not None else "",
        _UOM_LABEL.get(order.unit, order.unit) if order is not None else "",
        _show_num(Decimal(str(order.qty_order))) if order is not None else "",
        header,
        order.customer if order is not None else "",
        item.color if item is not None else "",
        _pack(names, ("气泡垫",)),
        _pack(names, ("真空袋", "内衬")),
        _pack(names, ("回料袋",)),
        "",
        _show_num(Decimal(str(sph.sph_value))) if sph is not None else "",
        task.crew_plan,
        _boxes(task.qty_board, board_per_box),
        "" if task.qty_actual is None else task.qty_actual,
        "",
        "",
        (wo.override_reason or "").strip(),
    ]
    return row, hours


def _pack(names: list[str], keys: tuple[str, ...]) -> str:
    hit = [name for name in names if any(key in name for key in keys)]
    return "、".join(hit)


def _boxes(qty_board: int, board_per_box: Decimal) -> str:
    if board_per_box <= 0:
        return ""
    return _show_num(Decimal(qty_board) / board_per_box)


def _hours(value: Decimal) -> str:
    return format(value.quantize(Decimal("0.01")), "f")


def _show_num(value: Decimal) -> str:
    text = format(value.quantize(Decimal("0.0001")), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def _sheet_date(value: date) -> str:
    return f"{value.year}-{value.month}-{value.day}"


def _finished_components(session: Session, wo: WoRow, orders: dict[str, SoOrderRow]) -> list[dict]:
    item = _load_item(session, wo.item_code)
    header = wo.source_order_no.split("#L")[0]
    order_row = orders.get(header)
    if item is None or order_row is None:
        return [{"code": "", "name": "无下级子件", "role": "", "gross": 0}]
    order = Order(
        order_no=header,
        customer=order_row.customer,
        item_code=wo.item_code,
        qty_order=order_row.qty_order,
        unit=Uom(order_row.unit),
        due_date=order_row.due_date,
        customer_level=order_row.customer_level,
    )
    lines = _bom_lines_for(session, wo.item_code)
    out: list[dict] = []
    if lines:
        for line in lines:
            gross = gross_board_for_line(order, line, item)
            child = session.get(MdItemRow, line.component_item_code)
            out.append(
                {
                    "code": line.component_item_code,
                    "name": child.item_name if child else line.component_item_code,
                    "role": role_label(line.component_role.value),
                    "gross": gross,
                }
            )
        return out
    route = _route_for(session, wo.item_code)
    if route and route.needs_semi and route.semi_item_code and route.semi_board_per_box:
        finished = _wo_model(wo)
        gross = gross_semi_board(finished, route, item)
        child = session.get(MdItemRow, route.semi_item_code)
        return [
            {
                "code": route.semi_item_code,
                "name": child.item_name if child else route.semi_item_code,
                "role": role_label("SEMI"),
                "gross": gross,
            }
        ]
    return [{"code": "", "name": "无下级子件", "role": "", "gross": 0}]


def _wo_model(wo: WoRow) -> Wo:
    return Wo(
        wo_no=wo.wo_no,
        wo_type=WoType(wo.wo_type),
        source_order_no=wo.source_order_no,
        item_code=wo.item_code,
        group_code=GroupCode(wo.group_code),
        dept=Dept(wo.dept),
        qty_order=wo.qty_order,
        qty_board_plan=wo.qty_board_plan,
        due_date=wo.due_date,
        earliest_start=wo.earliest_start,
        crew_plan=wo.crew_plan,
        status=WoStatus(wo.status),
    )
